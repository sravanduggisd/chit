import os
import shutil
import subprocess
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image

venkat_dir = r'd:\CHIT\03.VENKATESHWARA'
wb_path = os.path.join(venkat_dir, 'workbook VENKATESHWARA.xlsx')
backup_path = r'd:\CHIT\WORKSPACE\workbook_venkateshwara_backup.xlsx'

# Create safe backup
shutil.copyfile(wb_path, backup_path)
print(f"Backup created at: {backup_path}")

edge_exe = r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe'
images_dir = os.path.join(venkat_dir, 'Monthly Images')
os.makedirs(images_dir, exist_ok=True)

# 12-month schedule for Venkateshwara Chit (April 2026 to March 2027)
venkat_plan = [
    {
        'num': '01',
        'tab_name': 'April',
        'month_full': 'April',
        'year': '2026',
        'person': 'D DHANUNJAYA',
        'bid_amount': 23150,
        'payout': 262500,
    },
    {
        'num': '02',
        'tab_name': 'MAY',
        'month_full': 'May',
        'year': '2026',
        'person': 'Amara Pavan',
        'bid_amount': 22500,
        'payout': 262500,
    },
    {
        'num': '03',
        'tab_name': 'JUNE',
        'month_full': 'June',
        'year': '2026',
        'person': 'KIRAN (anil frd)',
        'bid_amount': 21900,
        'payout': 262500,
    },
    {
        'num': '04',
        'tab_name': 'JULY',
        'month_full': 'July',
        'year': '2026',
        'person': 'A RAKESH',
        'bid_amount': 21250,
        'payout': 262500,
    },
    {
        'num': '05',
        'tab_name': 'AUGUST',
        'month_full': 'August',
        'year': '2026',
        'person': 'BHASKAR',
        'bid_amount': 20650,
        'payout': 266250,
    },
    {
        'num': '06',
        'tab_name': 'SEPTEMBER',
        'month_full': 'September',
        'year': '2026',
        'person': 'P RANI',
        'bid_amount': 20000,
        'payout': 266250,
    },
    {
        'num': '07',
        'tab_name': 'OCTOBER',
        'month_full': 'October',
        'year': '2026',
        'person': 'SATISH',
        'bid_amount': 19400,
        'payout': 266250,
    },
    {
        'num': '08',
        'tab_name': 'NOVEMBER',
        'month_full': 'November',
        'year': '2026',
        'person': 'BOMBAY DYEING',
        'bid_amount': 18750,
        'payout': 268750,
    },
    {
        'num': '09',
        'tab_name': 'DECEMBER',
        'month_full': 'December',
        'year': '2026',
        'person': 'P HARI KRISHNA',
        'bid_amount': 18150,
        'payout': 268750,
    },
    {
        'num': '10',
        'tab_name': 'JANUARY',
        'month_full': 'January',
        'year': '2027',
        'person': 'CH RAJINI',
        'bid_amount': 17500,
        'payout': 272500,
    },
    {
        'num': '11',
        'tab_name': 'FEBRUARY',
        'month_full': 'February',
        'year': '2027',
        'person': 'MOUNIKA',
        'bid_amount': 16900,
        'payout': 275000,
    },
    {
        'num': '12',
        'tab_name': 'MARCH',
        'month_full': 'March',
        'year': '2027',
        'person': 'SWETHA',
        'bid_amount': 25000,
        'payout': 281250,
    }
]

all_members = [item['person'] for item in venkat_plan]

# ==========================================
# PART 1: UPDATE WORKBOOK VENKATESHWARA.XLSX
# ==========================================
wb = openpyxl.load_workbook(wb_path)

font_title = Font(name='Calibri', size=16, bold=True)
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
fill_navy = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')

font_taker = Font(name='Calibri', size=11, bold=True, color='000000')
fill_taker = PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid')

font_nontaker = Font(name='Calibri', size=11, bold=False, color='000000')
fill_nontaker = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

font_bold = Font(name='Calibri', size=11, bold=True)
font_regular = Font(name='Calibri', size=11)

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')

table_border = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000')
)
currency_fmt = '"₹"\\ #,##0;[Red]"₹"\\ \\-#,##0'

# Existing sheets to keep untouched: April, MAY, JUNE, JULY
sheets_to_update = ['AUGUST', 'SEPTEMBER', 'OCTOBER', 'NOVEMBER', 'DECEMBER', 'JANUARY', 'FEBRUARY', 'MARCH']

takers_cum = []
for idx, item in enumerate(venkat_plan):
    takers_cum.append(item['person'])
    tab = item['tab_name']
    
    if tab in sheets_to_update:
        if tab in wb.sheetnames:
            ws = wb[tab]
            for r in range(1, max(ws.max_row + 1, 30)):
                for c in range(1, max(ws.max_column + 1, 10)):
                    ws.cell(r, c).value = None
        else:
            ws = wb.create_sheet(title=tab)
            
        ws.column_dimensions['A'].width = 24.0
        ws.column_dimensions['B'].width = 28.0
        ws.column_dimensions['C'].width = 28.0
        
        # Title
        ws.merge_cells('A1:C1')
        ws['A1'] = f"                       VENKATESHWARA CHIT – {item['month_full'].upper()} MONTH {item['year']}"
        ws['A1'].font = font_title
        ws['A1'].alignment = align_center
        ws.row_dimensions[1].height = 28.0
        
        # Headers
        headers = [('A3', 'Person Name'), ('B3', 'Amount Need to Pay'), ('C3', 'paid status ')]
        for cell_ref, text in headers:
            cell = ws[cell_ref]
            cell.value = text
            cell.font = font_header
            cell.fill = fill_navy
            cell.alignment = align_center
            cell.border = table_border
        ws.row_dimensions[3].height = 24.0
        
        # Member rows
        row_idx = 4
        # Takers
        for name in takers_cum:
            c_name = ws.cell(row=row_idx, column=1, value=name)
            c_name.font = font_taker
            c_name.fill = fill_taker
            c_name.alignment = align_left
            c_name.border = table_border
            
            c_amt = ws.cell(row=row_idx, column=2, value='₹ 25,000')
            c_amt.font = font_taker
            c_amt.fill = fill_taker
            c_amt.alignment = align_center
            c_amt.number_format = currency_fmt
            c_amt.border = table_border
            
            c_status = ws.cell(row=row_idx, column=3, value=None)
            c_status.font = font_regular
            c_status.border = table_border
            
            ws.row_dimensions[row_idx].height = 20.0
            row_idx += 1
            
        # Non-Takers
        nontakers_now = [p for p in all_members if p not in takers_cum]
        for name in nontakers_now:
            c_name = ws.cell(row=row_idx, column=1, value=name)
            c_name.font = font_nontaker
            c_name.fill = fill_nontaker
            c_name.alignment = align_left
            c_name.border = table_border
            
            c_amt = ws.cell(row=row_idx, column=2, value=item['bid_amount'])
            c_amt.font = font_nontaker
            c_amt.fill = fill_nontaker
            c_amt.alignment = align_center
            c_amt.number_format = currency_fmt
            c_amt.border = table_border
            
            c_status = ws.cell(row=row_idx, column=3, value=None)
            c_status.font = font_regular
            c_status.border = table_border
            
            ws.row_dimensions[row_idx].height = 20.0
            row_idx += 1
            
        # Total Collection Row
        ws.cell(row=16, column=1, value="Total Collection").font = font_bold
        ws.cell(row=16, column=1).alignment = align_center
        ws.cell(row=16, column=1).border = table_border
        
        total_val = len(takers_cum) * 25000 + len(nontakers_now) * item['bid_amount']
        c_tot = ws.cell(row=16, column=2, value=f"₹ {total_val:,}")
        c_tot.font = font_bold
        c_tot.alignment = align_center
        c_tot.border = table_border
        
        ws.cell(row=16, column=3, value=None).border = table_border
        ws.row_dimensions[16].height = 22.0
        
        # Counting Section
        payout = item['payout']
        net_counting = payout - 25000
        
        ws.cell(row=19, column=2, value=f"Amount counting - {item['person'].upper()}").font = font_bold
        ws.cell(row=19, column=2).alignment = align_center
        
        counting_summary = f"Total - {item['tab_name']} {payout:,} REMOVE 25,000 (SELF PAID)   BALANCE {net_counting:,} COUNTING"
        ws.cell(row=20, column=1, value=counting_summary).font = font_bold
        
        ws.cell(row=22, column=1, value="ONLINE / UPI").font = font_regular
        ws.cell(row=23, column=1, value="CASH").font = font_regular
        ws.cell(row=24, column=1, value="PENDING").font = font_regular
        ws.cell(row=25, column=1, value="CHIT OWNER BALANCE").font = font_regular
        
        c_net = ws.cell(row=27, column=3, value=net_counting)
        c_net.font = font_bold
        c_net.number_format = currency_fmt

# Reorder sheets in chronological order
desired_order = ['April', 'MAY', 'JUNE', 'JULY', 'AUGUST', 'SEPTEMBER', 'OCTOBER', 'NOVEMBER', 'DECEMBER', 'JANUARY', 'FEBRUARY', 'MARCH']
wb._sheets = [wb[s] for s in desired_order if s in wb.sheetnames] + [wb[s] for s in wb.sheetnames if s not in desired_order]

wb.save(wb_path)
print(f"Venkateshwara Chit Workbook updated successfully: {wb_path}")


# ==========================================
# PART 2: CREATE ALL 12 ANNOUNCEMENT IMAGES
# ==========================================
html_template = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<style>
  * {
    margin: 0;
    padding: 0;
    box-sizing: border-box;
    font-family: Arial, "Helvetica Neue", Helvetica, sans-serif;
  }
  body {
    width: 1024px;
    height: 731px;
    background-color: #ffffff;
    margin: 0;
    padding: 0;
    display: flex;
    flex-direction: column;
    align-items: center;
    justify-content: flex-start;
  }
  .title-container {
    height: 110px;
    display: flex;
    align-items: center;
    justify-content: center;
    padding-top: 15px;
  }
  .title {
    font-size: 32px;
    font-weight: 800;
    color: #000000;
    letter-spacing: -0.5px;
    text-transform: uppercase;
  }
  .table-container {
    width: 934px;
  }
  table {
    width: 100%;
    border-collapse: collapse;
    border: 2px solid #000000;
  }
  th {
    background-color: #326ac7;
    color: #ffffff;
    font-size: 22px;
    font-weight: 700;
    height: 42px;
    border: 2px solid #000000;
    text-align: center;
    vertical-align: middle;
  }
  td {
    border: 2px solid #000000;
    height: 42px;
    font-size: 21px;
    font-weight: 700;
    vertical-align: middle;
  }
  .col-name {
    width: 340px;
    text-align: left;
    padding-left: 14px;
  }
  .col-amt {
    width: 594px;
    text-align: right;
    padding-right: 20px;
  }
  .taker {
    background-color: #f70305;
    color: #ffffff;
  }
  .nontaker {
    background-color: #9be278;
    color: #000000;
  }
</style>
</head>
<body>
  <div class="title-container">
    <div class="title"><!--TITLE--></div>
  </div>
  <div class="table-container">
    <table>
      <thead>
        <tr>
          <th class="col-name" style="text-align: center;">Person Name</th>
          <th class="col-amt" style="text-align: center;">Amount Need to Pay</th>
        </tr>
      </thead>
      <tbody>
        <!--ROWS-->
      </tbody>
    </table>
  </div>
</body>
</html>
"""

takers_cum_img = []
for item in venkat_plan:
    takers_cum_img.append(item['person'])
    nontakers_cum_img = [p for p in all_members if p not in takers_cum_img]
    
    title = f"VENKATESHWARA CHIT – {item['month_full'].upper()} MONTH"
    file_name = f"{item['num']}. {item['month_full'].upper()} - {item['year']}.jpg"
    
    rows_html = []
    for name in takers_cum_img:
        rows_html.append(f"""
        <tr class="taker">
          <td class="col-name">{name}</td>
          <td class="col-amt">25,000</td>
        </tr>
        """)
        
    for name in nontakers_cum_img:
        rows_html.append(f"""
        <tr class="nontaker">
          <td class="col-name">{name}</td>
          <td class="col-amt">{item['bid_amount']:,}</td>
        </tr>
        """)
        
    html_content = html_template.replace('<!--TITLE-->', title).replace('<!--ROWS-->', "".join(rows_html))
    
    temp_html = os.path.join(r'd:\CHIT\WORKSPACE', f"temp_venkat_{item['num']}.html")
    with open(temp_html, 'w', encoding='utf-8') as f:
        f.write(html_content)
        
    temp_png = os.path.join(r'd:\CHIT\WORKSPACE', f"temp_venkat_{item['num']}.png")
    final_jpg = os.path.join(images_dir, file_name)
    
    cmd = [
        edge_exe,
        '--headless',
        '--disable-gpu',
        '--force-device-scale-factor=1',
        f'--screenshot={temp_png}',
        '--window-size=1024,731',
        f'file:///{temp_html.replace(os.sep, "/")}'
    ]
    subprocess.run(cmd, check=True)
    
    img = Image.open(temp_png)
    img.convert('RGB').save(final_jpg, 'JPEG', quality=95)
    print(f"Generated {file_name}")
    
    if os.path.exists(temp_html):
        os.remove(temp_html)
    if os.path.exists(temp_png):
        os.remove(temp_png)

print("All 12 Venkateshwara Chit images generated successfully!")
