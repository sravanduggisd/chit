import openpyxl

wb = openpyxl.load_workbook(r'd:\CHIT\01.GANAPATHI\Ganapathi_Chit_April- WORKBOOK.xlsx')
with open(r'd:\CHIT\WORKSPACE\verify_output.txt', 'w', encoding='utf-8') as f:
    f.write(f"Sheet names: {wb.sheetnames}\n\n")
    for sname in ['SEP', 'OCT', 'NOV', 'DEC', 'JAN', 'FEB']:
        ws = wb[sname]
        f.write(f"*** TAB: {sname} ***\n")
        f.write(f"Title: {ws['A1'].value.strip()}\n")
        for r in range(4, 17):
            f.write(f"  R{r:2d}: {ws.cell(r, 1).value} | {ws.cell(r, 2).value}\n")
        f.write(f"Counting header: {ws.cell(19, 2).value}\n")
        f.write(f"Counting summary: {ws.cell(20, 1).value}\n")
        f.write(f"Row 27 Net: {ws.cell(27, 3).value}\n\n")

print("Done")
