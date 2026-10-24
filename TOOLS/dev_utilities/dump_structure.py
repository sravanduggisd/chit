import openpyxl
import json

wb = openpyxl.load_workbook(r'd:\CHIT\01.GANAPATHI\Ganapathi_Chit_April- WORKBOOK.xlsx', data_only=False)

data = {}
for name in wb.sheetnames:
    ws = wb[name]
    sheet_info = {
        'max_row': ws.max_row,
        'max_col': ws.max_column,
        'merged_cells': [str(m) for m in ws.merged_cells.ranges],
        'rows': []
    }
    for r in range(1, ws.max_row + 1):
        row_vals = []
        has_val = False
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            val = cell.value
            if val is not None:
                has_val = True
            row_vals.append({
                'col': c,
                'val': str(val) if val is not None else None,
                'number_format': cell.number_format,
                'font_bold': cell.font.bold if cell.font else False,
                'font_size': cell.font.size if cell.font else None,
                'font_name': cell.font.name if cell.font else None,
            })
        if has_val:
            sheet_info['rows'].append({'row': r, 'cells': row_vals})
    data[name] = sheet_info

with open(r'd:\CHIT\WORKSPACE\ganapathi_sheet_analysis.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, indent=2, ensure_ascii=False)

print("Dump completed successfully.")
