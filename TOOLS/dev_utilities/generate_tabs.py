import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import shutil

src_file = r'd:\CHIT\01.GANAPATHI\Ganapathi_Chit_April- WORKBOOK.xlsx'
backup_file = r'd:\CHIT\WORKSPACE\Ganapathi_Chit_April- WORKBOOK_backup.xlsx'

# 1. Create a safe backup
shutil.copyfile(src_file, backup_file)
print(f"Backup created at: {backup_file}")

wb = openpyxl.load_workbook(src_file)

# Common styles
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
table_border = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000')
)

font_title = Font(name='Calibri', size=16, bold=True)
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
fill_header = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')

# Red fill for chit takers
font_taker = Font(name='Calibri', size=11, bold=True, color='000000')
fill_taker = PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid')

# Light green fill for non-takers
font_nontaker = Font(name='Calibri', size=11, bold=False, color='000000')
fill_nontaker = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

font_bold = Font(name='Calibri', size=11, bold=True)
font_regular = Font(name='Calibri', size=11)

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

currency_fmt = '"₹"\\ #,##0;[Red]"₹"\\ \\-#,##0'

# Schedule of remaining months based on Ganapathi.png & existing Aug sheet:
# In Aug, the 6 takers are: Sarath Babu, Anil, Jyothi, Girl (Giri), Anjineyulu, Ramanaiah
# Remaining to take: Harish, Pallavi, P. Hari, Sumanth, Vamsi Madhu, Uma Suri
remaining_months = [
    {
        'tab_name': 'SEP',
        'title': 'Ganapathi Chit – SEP Month 2026',
        'taker_name': 'Harish',
        'payout': 213000,
        'bid_amount': 15500,
        'takers_order': ['Sarath Babu', 'Anil', 'Jyothi', 'Girl', 'Anjineyulu', 'Ramanaiah', 'Harish'],
        'nontakers_order': ['Pallavi', 'P. Hari', 'Sumanth', 'Vamsi Madhu', 'Uma Suri']
    },
    {
        'tab_name': 'OCT',
        'title': 'Ganapathi Chit – OCT Month 2026',
        'taker_name': 'Pallavi',
        'payout': 215000,
        'bid_amount': 15000,
        'takers_order': ['Sarath Babu', 'Anil', 'Jyothi', 'Girl', 'Anjineyulu', 'Ramanaiah', 'Harish', 'Pallavi'],
        'nontakers_order': ['P. Hari', 'Sumanth', 'Vamsi Madhu', 'Uma Suri']
    },
    {
        'tab_name': 'NOV',
        'title': 'Ganapathi Chit – NOV Month 2026',
        'taker_name': 'P. Hari',
        'payout': 215000,
        'bid_amount': 14500,
        'takers_order': ['Sarath Babu', 'Anil', 'Jyothi', 'Girl', 'Anjineyulu', 'Ramanaiah', 'Harish', 'Pallavi', 'P. Hari'],
        'nontakers_order': ['Sumanth', 'Vamsi Madhu', 'Uma Suri']
    },
    {
        'tab_name': 'DEC',
        'title': 'Ganapathi Chit – DEC Month 2026',
        'taker_name': 'Sumanth',
        'payout': 218000,
        'bid_amount': 14000,
        'takers_order': ['Sarath Babu', 'Anil', 'Jyothi', 'Girl', 'Anjineyulu', 'Ramanaiah', 'Harish', 'Pallavi', 'P. Hari', 'Sumanth'],
        'nontakers_order': ['Vamsi Madhu', 'Uma Suri']
    },
    {
        'tab_name': 'JAN',
        'title': 'Ganapathi Chit – JAN Month 2027',
        'taker_name': 'Vamsi Madhu',
        'payout': 220000,
        'bid_amount': 13500,
        'takers_order': ['Sarath Babu', 'Anil', 'Jyothi', 'Girl', 'Anjineyulu', 'Ramanaiah', 'Harish', 'Pallavi', 'P. Hari', 'Sumanth', 'Vamsi Madhu'],
        'nontakers_order': ['Uma Suri']
    },
    {
        'tab_name': 'FEB',
        'title': 'Ganapathi Chit – FEB Month 2027',
        'taker_name': 'Uma Suri',
        'payout': 225000,
        'bid_amount': 20000,
        'takers_order': ['Sarath Babu', 'Anil', 'Jyothi', 'Girl', 'Anjineyulu', 'Ramanaiah', 'Harish', 'Pallavi', 'P. Hari', 'Sumanth', 'Vamsi Madhu', 'Uma Suri'],
        'nontakers_order': []
    }
]

for m in remaining_months:
    tab = m['tab_name']
    if tab in wb.sheetnames:
        ws = wb[tab]
    else:
        ws = wb.create_sheet(title=tab)
    
    # Column dimensions
    ws.column_dimensions['A'].width = 22.0
    ws.column_dimensions['B'].width = 32.0
    ws.column_dimensions['C'].width = 28.0
    ws.column_dimensions['D'].width = 16.0
    
    # Row 1: Title
    ws.merge_cells('A1:C1')
    ws['A1'] = f"                       {m['title']}"
    ws['A1'].font = font_title
    ws['A1'].alignment = align_center
    ws.row_dimensions[1].height = 28.0
    
    # Row 3: Headers
    headers = [('A3', 'Name'), ('B3', 'Amount Need to Pay'), ('C3', 'paid status ')]
    for cell_ref, text in headers:
        cell = ws[cell_ref]
        cell.value = text
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = table_border
    ws.row_dimensions[3].height = 24.0
    
    # Rows 4-15: Members
    row_idx = 4
    # 1. Takers (Red fill, 20,000)
    for name in m['takers_order']:
        c_name = ws.cell(row=row_idx, column=1, value=name)
        c_name.font = font_taker
        c_name.fill = fill_taker
        c_name.alignment = align_left
        c_name.border = table_border
        
        c_amt = ws.cell(row=row_idx, column=2, value='₹ 20,000')
        c_amt.font = font_taker
        c_amt.fill = fill_taker
        c_amt.alignment = align_center
        c_amt.number_format = currency_fmt
        c_amt.border = table_border
        
        c_status = ws.cell(row=row_idx, column=3, value=None)
        c_status.font = font_regular
        c_status.border = table_border
        
        ws.row_dimensions[row_idx].height = 20.0
        row_idx += 1
        
    # 2. Non-Takers (Light green fill, bid_amount)
    for name in m['nontakers_order']:
        c_name = ws.cell(row=row_idx, column=1, value=name)
        c_name.font = font_nontaker
        c_name.fill = fill_nontaker
        c_name.alignment = align_left
        c_name.border = table_border
        
        c_amt = ws.cell(row=row_idx, column=2, value=m['bid_amount'])
        c_amt.font = font_nontaker
        c_amt.fill = fill_nontaker
        c_amt.alignment = align_center
        c_amt.number_format = currency_fmt
        c_amt.border = table_border
        
        c_status = ws.cell(row=row_idx, column=3, value=None)
        c_status.font = font_regular
        c_status.border = table_border
        
        ws.row_dimensions[row_idx].height = 20.0
        row_idx += 1
        
    # Row 16: Total
    ws.cell(row=16, column=1, value="Total Collection").font = font_bold
    ws.cell(row=16, column=1).alignment = align_center
    ws.cell(row=16, column=1).border = table_border
    
    total_val = len(m['takers_order']) * 20000 + len(m['nontakers_order']) * m['bid_amount']
    c_tot = ws.cell(row=16, column=2, value=f"₹ {total_val:,}")
    c_tot.font = font_bold
    c_tot.alignment = align_center
    c_tot.border = table_border
    
    ws.cell(row=16, column=3, value=None).border = table_border
    ws.row_dimensions[16].height = 22.0
    
    # Counting Section (like July / May sheets)
    payout = m['payout']
    net_counting = payout - 20000
    
    ws.cell(row=19, column=2, value=f"Amount counting - {m['taker_name'].upper()}").font = font_bold
    ws.cell(row=19, column=2).alignment = align_center
    
    counting_summary = f"Total - {m['tab_name']} {payout:,} REMOVE 20,000 (SELF PAID)   BALANCE {net_counting:,} COUNTING"
    ws.cell(row=20, column=1, value=counting_summary).font = font_bold
    
    ws.cell(row=22, column=1, value="ONLINE / UPI").font = font_regular
    ws.cell(row=22, column=2, value="")
    ws.cell(row=22, column=3, value="")
    
    ws.cell(row=23, column=1, value="CASH").font = font_regular
    ws.cell(row=23, column=2, value="")
    ws.cell(row=23, column=3, value="")
    
    ws.cell(row=24, column=1, value="PENDING").font = font_regular
    ws.cell(row=24, column=2, value="")
    ws.cell(row=24, column=3, value="")
    
    ws.cell(row=25, column=1, value="CHIT OWNER BALANCE").font = font_regular
    ws.cell(row=25, column=2, value="")
    ws.cell(row=25, column=3, value="")
    
    c_net = ws.cell(row=27, column=3, value=net_counting)
    c_net.font = font_bold
    c_net.number_format = currency_fmt

wb.save(src_file)
print("Workbook updated successfully with all remaining tabs!")
