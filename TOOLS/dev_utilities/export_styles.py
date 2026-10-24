import openpyxl

wb = openpyxl.load_workbook(r'd:\CHIT\01.GANAPATHI\Ganapathi_Chit_April- WORKBOOK.xlsx')

with open(r'd:\CHIT\WORKSPACE\cell_styles_detailed.txt', 'w', encoding='utf-8') as out:
    for sheetname in ['April', 'JULY', 'AUG']:
        ws = wb[sheetname]
        out.write(f"================ SHEET: {sheetname} ================\n")
        for r in range(1, 17):
            c1 = ws.cell(r, 1)
            c2 = ws.cell(r, 2)
            c3 = ws.cell(r, 3)
            
            def info(cell):
                val = cell.value
                font = f"{cell.font.name} {cell.font.size} b={cell.font.bold} c={cell.font.color.rgb if cell.font and cell.font.color else None}"
                fill_color = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None
                fill_type = cell.fill.fill_type if cell.fill else None
                align = f"h={cell.alignment.horizontal} v={cell.alignment.vertical}" if cell.alignment else ""
                border = {s: getattr(cell.border, s).style for s in ['top', 'bottom', 'left', 'right']} if cell.border else {}
                return f"val={val!r} | font={font} | fill={fill_type}:{fill_color} | align={align} | num_fmt={cell.number_format}"

            out.write(f"R{r:2d} C1: {info(c1)}\n")
            out.write(f"    C2: {info(c2)}\n")
            out.write(f"    C3: {info(c3)}\n")

print("Styles exported.")
