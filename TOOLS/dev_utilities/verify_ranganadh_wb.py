import openpyxl

wb = openpyxl.load_workbook(r'd:\CHIT\04.Ranganadh Chit\Ranganadh_Chit_WORKBOOK.xlsx')
with open(r'd:\CHIT\WORKSPACE\ranganadh_verify.txt', 'w', encoding='utf-8') as f:
    f.write(f"Sheets: {wb.sheetnames}\n\n")
    for s in wb.sheetnames:
        ws = wb[s]
        f.write(f"*** SHEET: {s} ***\n")
        f.write(f"Title: {ws['A1'].value.strip() if ws['A1'].value else ''}\n")
        f.write(f"Row 4: {ws.cell(4, 1).value} | {ws.cell(4, 2).value}\n")
        f.write(f"Row 16: {ws.cell(16, 1).value} | {ws.cell(16, 2).value}\n")
        if s != 'Master Chit Tracker':
            f.write(f"Counting: {ws.cell(19, 2).value}\n")
            f.write(f"Summary: {ws.cell(20, 1).value}\n")
            f.write(f"Net counting: {ws.cell(27, 3).value}\n")
        f.write("\n")

print("Done verifying Ranganadh workbook.")
